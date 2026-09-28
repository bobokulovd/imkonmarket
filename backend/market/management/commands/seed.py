"""Katalogni yuklash va har bir muassasaga login/parol ochish.

    python manage.py seed                       # katalog + loginlar (mavjud parollar o'zgarmaydi)
    python manage.py seed --reset-passwords     # hammaga yangi parol
    python manage.py seed --credentials out.xlsx
"""
import secrets
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand
from django.db import transaction

from market.i18n_meta import catalog
from market.models import Category, Product, Profile, Seller

ALPH = "ABCDEFGHJKMNPQRSTUVWXYZabcdefghjkmnpqrstuvwxyz23456789"


def gen_password():
    core = "".join(secrets.choice(ALPH) for _ in range(9))
    return f"{core[:4]}-{core[4:]}"


class Command(BaseCommand):
    help = "Katalog, muassasalar va loginlarni yuklaydi"

    def add_arguments(self, parser):
        parser.add_argument("--reset-passwords", action="store_true")
        parser.add_argument("--credentials", default=str(Path(settings.BASE_DIR) / "credentials.xlsx"))
        parser.add_argument("--no-products", action="store_true")
        parser.add_argument("--if-empty", action="store_true", help="Baza bo'sh bo'lsagina ishlaydi (konteyner start uchun)")
        parser.add_argument("--passwords-from", help="Loginlar Excel faylidan parollarni o'rnatish (Login/Parol ustunlari)")

    @transaction.atomic
    def handle(self, *args, **opts):
        if opts.get("if_empty") and Seller.objects.exists():
            self.stdout.write("Baza to'ldirilgan — seed o'tkazib yuborildi")
            return
        data = catalog()
        User = get_user_model()
        for i, c in enumerate(data["categories"]):
            Category.objects.update_or_create(slug=c["slug"], defaults={"icon": c["icon"], "name": c["name"], "order": i})
        cats = {c.slug: c for c in Category.objects.all()}

        preset = {}
        import json
        import os
        if os.environ.get("SEED_PASSWORDS"):  # {"jiek14": "parol", ...} — Dokploy environment orqali
            preset.update(json.loads(os.environ["SEED_PASSWORDS"]))
        if opts.get("passwords_from"):
            from openpyxl import load_workbook
            ws = load_workbook(opts["passwords_from"], read_only=True)["Loginlar"]
            rows = list(ws.iter_rows(values_only=True))
            hi = {h: i for i, h in enumerate(rows[0])}
            for r in rows[1:]:
                if r and r[hi["Login"]] and r[hi["Parol"]] and not str(r[hi["Parol"]]).startswith("("):
                    preset[str(r[hi["Login"]])] = str(r[hi["Parol"]])
        creds = []
        sellers = {}
        for s in data["sellers"]:
            obj, _ = Seller.objects.update_or_create(code=s["code"], defaults={
                "role": s["role"], "name": s["name"], "inn": s["inn"], "region": s["region"],
                "region_i18n": s["region_i18n"], "district_i18n": s["district_i18n"],
            })
            sellers[s["code"]] = obj
            username = s["code"].replace("-", "")
            user = User.objects.filter(username=username).first()
            password = None
            if not user:
                password = preset.get(username) or gen_password()
                user = User.objects.create_user(username=username, password=password)
            elif opts["reset_passwords"] or username in preset:
                password = preset.get(username) or gen_password()
                user.set_password(password)
            if s["role"] == "operator":
                user.is_staff = True
                user.is_superuser = True
            user.save()
            Profile.objects.update_or_create(user=user, defaults={"seller": obj})
            creds.append((obj, s, username, password))

        n_new = 0
        if not opts["no_products"]:
            for p in data["products"]:
                obj = Product.objects.filter(sku=p["sku"]).first()
                fields = {
                    "seller": sellers[p["seller"]], "category": cats[p["category"]], "name": p["name"],
                    "spec": p["spec"], "unit": p["unit"], "price": p["price"], "daily_capacity": p["daily_capacity"],
                    "delivery": p["delivery"], "address": p["address"], "note": p["note"],
                }
                if obj is None:
                    obj = Product(sku=p["sku"], stock=p["stock"], **fields)
                    n_new += 1
                else:  # qoldiq va narxni sotuvchi o'zgartirgan bo'lishi mumkin — faqat tavsif yangilanadi
                    for k in ("name", "spec", "unit", "address", "category"):
                        setattr(obj, k, fields[k])
                obj.save()

        self.stdout.write(self.style.SUCCESS(
            f"Kategoriyalar: {len(cats)}, muassasalar: {len(sellers)}, yangi mahsulotlar: {n_new}, jami: {Product.objects.count()}"))
        if any(c[3] for c in creds):
            self._write_credentials(opts["credentials"], creds, data)

    def _write_credentials(self, path, creds, data):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Loginlar"
        site = settings.SITE_URL.rstrip("/")
        head = ["№", "Muassasa", "Hudud", "Tuman/shahar", "STIR", "Login", "Parol", "Mahsulotlar", "Kabinet", "Rol"]
        ws.append(head)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        counts = {}
        for p in data["products"]:
            counts[p["seller"]] = counts.get(p["seller"], 0) + 1
        for i, (obj, s, username, password) in enumerate(creds, 1):
            ws.append([
                i, s["name"]["uz"], s["region"], s["district_i18n"]["uz"], s["inn"], username,
                password or "(o'zgarmagan)", counts.get(s["code"], 0), f"{site}/cabinet",
                "Operator (admin)" if s["role"] == "operator" else "Sotuvchi",
            ])
        for col, w in zip("ABCDEFGHIJ", [5, 16, 28, 24, 12, 12, 14, 12, 34, 16]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

        q = wb.create_sheet("Ma'lumot sifati")
        q.append(["SKU", "Muassasa", "Mahsulot", "Xususiyat", "Izoh"])
        for cell in q[1]:
            cell.font = Font(bold=True)
        for n in data["quality_notes"]:
            q.append([n["sku"], n["seller"], n["name"], n["spec"], n["note"]])
        extra = [
            ("—", "jiek-02", "Fayl: 2-JIEK", "", "Excel sarlavhasida «4-son ЖИЭК» deb yozilgan, manzil Qarshi — 2-son JIEK ga biriktirildi"),
            ("—", "jiek-11", "Fayl: 11-son", "", "«Бир бирлик нархи» ustunida jami summa yozilgan — birlik narx = summa ÷ miqdor qilib hisoblandi"),
            ("—", "jiek-11", "Muassasa manzili", "", "Ro'yxatda Navoiy shahri, mahsulot faylida Karmana tumani — Karmana olindi"),
            ("—", "mk-44", "Muassasa manzili", "", "Ro'yxatda Ohangaron shahri, mahsulot faylida Zangiota tumani — Zangiota olindi"),
            ("—", "mk-29", "Barcha mahsulotlar", "", "Narx ustunida 5 so'm — narx «kelishiladi» deb qo'yildi, sotuvchi kabinetdan kiritadi"),
            ("—", "barchasi", "Qoldiq", "", "Excelda qoldiq yo'q — boshlang'ich qoldiq sifatida kunlik ishlab chiqarish quvvati qo'yildi; muassasalar kabinetdan yangilaydi"),
        ]
        for r in extra:
            q.append(list(r))
        for col, w in zip("ABCDE", [14, 10, 40, 24, 90]):
            q.column_dimensions[col].width = w
        for row in q.iter_rows(min_row=2):
            row[4].alignment = Alignment(wrap_text=True)
        wb.save(path)
        self.stdout.write(self.style.WARNING(f"Loginlar fayli: {path}"))
